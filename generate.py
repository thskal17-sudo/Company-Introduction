#!/usr/bin/env python3
"""회사 소개서 자동 생성기.

data/company.yaml 을 읽어 HTML / PDF / DOCX 소개서를 output/ 에 만듭니다.

사용법:
    python generate.py              # HTML 만 생성
    python generate.py --all        # HTML + PDF + DOCX
    python generate.py --pdf --docx # 원하는 형식만 추가
    python generate.py -d data/other.yaml -o dist
"""
from __future__ import annotations

import argparse
import base64
import datetime as dt
import mimetypes
import os
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import yaml
from jinja2 import Environment, FileSystemLoader, select_autoescape

ROOT = Path(__file__).resolve().parent
DEFAULT_DATA = ROOT / "data" / "company.yaml"
TEMPLATE_DIR = ROOT / "templates"
DEFAULT_OUT = ROOT / "output"


# --------------------------------------------------------------------------- #
# 데이터 로딩 / 정규화
# --------------------------------------------------------------------------- #
def load_data(path: Path) -> dict:
    with path.open(encoding="utf-8") as f:
        data = yaml.safe_load(f) or {}
    if not isinstance(data, dict) or not data.get("company", {}).get("name"):
        sys.exit(f"[오류] {path} 에 company.name 이 필요합니다.")
    return data


def paragraphs(text) -> list[str]:
    """여러 줄 문자열을 빈 줄 기준 문단 리스트로 변환."""
    if not text:
        return []
    if isinstance(text, list):
        return [str(t).strip() for t in text if str(t).strip()]
    blocks = str(text).strip().split("\n\n")
    return [" ".join(line.strip() for line in b.splitlines()).strip() for b in blocks if b.strip()]


def image_data_uri(rel_path: str, base: Path) -> str:
    """이미지 파일을 data URI 로 변환 (HTML 단일 파일용). 없으면 빈 문자열."""
    if not rel_path:
        return ""
    p = Path(rel_path)
    if not p.is_absolute():
        p = base / p
    if not p.is_file():
        print(f"[경고] 이미지를 찾을 수 없습니다: {rel_path}", file=sys.stderr)
        return ""
    mime = mimetypes.guess_type(p.name)[0] or "image/png"
    return f"data:{mime};base64,{base64.b64encode(p.read_bytes()).decode()}"


def normalize(raw: dict, base: Path) -> dict:
    d = dict(raw)
    c = dict(d.get("company") or {})
    c["logo_path"] = c.get("logo") or ""
    c["logo"] = image_data_uri(c.get("logo") or "", base)
    c["logo_white"] = image_data_uri(c.get("logo_white") or "", base)
    d["c"] = c
    d["about"] = paragraphs(d.get("about"))
    greeting = dict(d.get("greeting") or {})
    greeting["text"] = paragraphs(greeting.get("text"))
    d["greeting"] = greeting
    d["programs_heading"] = d.get("programs_heading") or {}
    instructors = []
    for i in d.get("instructors") or []:
        i = dict(i)
        i["photo_path"] = i.get("photo") or ""
        i["photo"] = image_data_uri(i.get("photo") or "", base)
        i["bio"] = paragraphs(i.get("bio"))
        instructors.append(i)
    d["instructors"] = instructors
    for key in ("values", "strategies", "stats", "history", "programs", "strengths", "clients", "testimonials", "program_axes"):
        d[key] = d.get(key) or []
    d["comparison"] = d.get("comparison") or {}
    d["program_axes_note"] = d.get("program_axes_note") or ""
    programs = []
    for p in d["programs"]:
        p = dict(p)
        p["sections"] = p.get("sections") or []
        p["photo_paths"] = p.get("photos") or []
        p["photos"] = [u for u in (image_data_uri(ph, base) for ph in p["photo_paths"]) if u]
        programs.append(p)
    d["programs"] = programs
    d["closing"] = d.get("closing") or {}
    d["generated_on"] = dt.date.today().strftime("%Y. %m")
    return d


# --------------------------------------------------------------------------- #
# HTML
# --------------------------------------------------------------------------- #
def emph(text) -> "Markup":
    """**강조** 표기를 <b>강조</b> 로 바꿉니다 (나머지 텍스트는 HTML 이스케이프)."""
    import re
    from markupsafe import Markup, escape
    parts = re.split(r"\*\*(.+?)\*\*", str(text or ""))
    out = []
    for i, part in enumerate(parts):
        out.append(f"<b>{escape(part)}</b>" if i % 2 else str(escape(part)))
    return Markup("".join(out))


def build_html(data: dict, out: Path) -> Path:
    env = Environment(
        loader=FileSystemLoader(str(TEMPLATE_DIR)),
        autoescape=select_autoescape(["html", "j2"]),
        trim_blocks=True,
        lstrip_blocks=True,
    )
    env.filters["emph"] = emph
    html = env.get_template("brochure.html.j2").render(**data)
    out.write_text(html, encoding="utf-8")
    return out


# --------------------------------------------------------------------------- #
# PDF (헤드리스 Chromium)
# --------------------------------------------------------------------------- #
def find_chrome() -> str | None:
    for env_key in ("CHROME_PATH", "CHROMIUM_PATH"):
        if os.environ.get(env_key) and Path(os.environ[env_key]).exists():
            return os.environ[env_key]
    for name in ("google-chrome", "google-chrome-stable", "chromium", "chromium-browser", "chrome", "msedge"):
        p = shutil.which(name)
        if p:
            return p
    pw = Path(os.environ.get("PLAYWRIGHT_BROWSERS_PATH", Path.home() / ".cache" / "ms-playwright"))
    for cand in sorted(pw.glob("chromium*/chrome-linux/chrome"), reverse=True):
        return str(cand)
    for cand in sorted(pw.glob("chromium*/chrome-mac/Chromium.app/Contents/MacOS/Chromium"), reverse=True):
        return str(cand)
    mac = "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"
    if Path(mac).exists():
        return mac
    for win in (r"C:\Program Files\Google\Chrome\Application\chrome.exe",
                r"C:\Program Files (x86)\Google\Chrome\Application\chrome.exe"):
        if Path(win).exists():
            return win
    return None


def build_pdf(html_path: Path, out: Path) -> Path | None:
    chrome = find_chrome()
    if not chrome:
        print("[경고] Chrome/Chromium 을 찾지 못해 PDF 를 건너뜁니다. "
              "CHROME_PATH 환경변수로 경로를 지정하거나 브라우저에서 HTML 을 열어 '인쇄 > PDF 저장' 하세요.",
              file=sys.stderr)
        return None
    with tempfile.TemporaryDirectory() as profile:
        cmd = [
            chrome, "--headless=new", "--disable-gpu", "--no-sandbox",
            f"--user-data-dir={profile}", "--no-pdf-header-footer",
            "--run-all-compositor-stages-before-draw", "--virtual-time-budget=8000",
            f"--print-to-pdf={out}", html_path.resolve().as_uri(),
        ]
        r = subprocess.run(cmd, capture_output=True, text=True)
    if r.returncode != 0 or not out.exists():
        print(f"[경고] PDF 생성 실패:\n{r.stderr[-800:]}", file=sys.stderr)
        return None
    check_overflow(html_path, out)
    return out


def check_overflow(html_path: Path, pdf_path: Path) -> None:
    """섹션 수보다 PDF 페이지가 많으면 어떤 섹션이 A4 한 장을 넘겼다는 뜻이므로 경고."""
    import re
    sections = html_path.read_text(encoding="utf-8").count('<section class="page')
    pages = len(re.findall(rb"/Type\s*/Page[^s]", pdf_path.read_bytes()))
    if pages > sections:
        print(f"[경고] 섹션 {sections}개가 PDF {pages}페이지로 출력되었습니다. "
              "어떤 섹션의 내용이 A4 한 장을 넘겼습니다. 해당 섹션의 문장을 줄이거나 항목 수를 줄여 주세요.",
              file=sys.stderr)


# --------------------------------------------------------------------------- #
# DOCX (python-docx)
# --------------------------------------------------------------------------- #
def build_docx(data: dict, out: Path, base: Path) -> Path | None:
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml.ns import qn
        from docx.shared import Cm, Pt, RGBColor
    except ImportError:
        print("[경고] python-docx 가 없어 DOCX 를 건너뜁니다. `pip install python-docx`", file=sys.stderr)
        return None

    c = data["c"]
    brand = RGBColor.from_string((c.get("brand_color") or "#1F4E79").lstrip("#"))

    doc = Document()
    for section in doc.sections:
        section.left_margin = section.right_margin = Cm(2)
        section.top_margin = section.bottom_margin = Cm(2)

    # 기본 한글 폰트 설정
    style = doc.styles["Normal"]
    style.font.name = "맑은 고딕"
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    for name in ("Heading 1", "Heading 2", "Heading 3", "Title"):
        st = doc.styles[name]
        st.font.name = "맑은 고딕"
        st.font.color.rgb = brand
        rpr = st.element.get_or_add_rPr()
        rpr.get_or_add_rFonts().set(qn("w:eastAsia"), "맑은 고딕")

    def heading(text, level=1):
        return doc.add_heading(text, level=level)

    def bullets(items):
        for it in items or []:
            doc.add_paragraph(str(it), style="List Bullet")

    def kv_table(rows):
        rows = [(k, v) for k, v in rows if v]
        if not rows:
            return
        t = doc.add_table(rows=0, cols=2)
        t.style = "Light List Accent 1"
        for k, v in rows:
            r = t.add_row().cells
            r[0].text, r[1].text = k, str(v)
            r[0].width, r[1].width = Cm(4), Cm(13)

    # ---- 표지
    if c.get("logo_path") and (base / c["logo_path"]).is_file():
        doc.add_picture(str(base / c["logo_path"]), width=Cm(8))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for _ in range(6):
        doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("회사소개서"); r.font.size = Pt(12); r.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(c["name"]); r.bold = True; r.font.size = Pt(34); r.font.color.rgb = brand
    if c.get("name_en"):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(c["name_en"]).font.size = Pt(12)
    if c.get("slogan"):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(c["slogan"]).font.size = Pt(16)
    if c.get("tagline"):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(c["tagline"]).font.size = Pt(11)
    for _ in range(10):
        doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(" | ".join(x for x in (c.get("website"), c.get("email"), c.get("phone")) if x)).font.size = Pt(10)
    doc.add_page_break()

    # ---- 회사 소개
    heading("회사 소개", 1)
    if data["greeting"]["text"]:
        heading("인사말", 2)
        for para in data["greeting"]["text"]:
            doc.add_paragraph(para)
        if data["greeting"].get("signer"):
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p.add_run(data["greeting"]["signer"]).bold = True
    for para in data["about"]:
        doc.add_paragraph(para)
    if data["stats"]:
        t = doc.add_table(rows=2, cols=len(data["stats"]))
        t.style = "Light Grid Accent 1"
        for i, s in enumerate(data["stats"]):
            cell = t.rows[0].cells[i]; cell.text = ""
            run = cell.paragraphs[0].add_run(str(s.get("value", "")))
            run.bold = True; run.font.size = Pt(18); run.font.color.rgb = brand
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            t.rows[1].cells[i].text = str(s.get("label", ""))
            t.rows[1].cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_paragraph()
    if data.get("mission"):
        p = doc.add_paragraph(); p.add_run("미션  ").bold = True; p.add_run(data["mission"])
    if data.get("vision"):
        p = doc.add_paragraph(); p.add_run("비전  ").bold = True; p.add_run(data["vision"])
    kv_table([("설립", c.get("founded")), ("대표", c.get("ceo")), ("주소", c.get("address")),
              ("사업자등록번호", c.get("business_number"))])

    # ---- 핵심 가치
    if data["values"]:
        heading("핵심 가치", 1)
        for i, v in enumerate(data["values"], 1):
            heading(f"{i}. {v.get('title', '')}", 3)
            if v.get("description"):
                doc.add_paragraph(v["description"])

    # ---- 전략 방향
    if data["strategies"]:
        heading("전략 방향", 1)
        for st in data["strategies"]:
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(str(st.get("title", ""))).bold = True
            if st.get("description"):
                p.add_run(f" — {st['description']}")

    # ---- 연혁
    if data["history"]:
        heading("연혁", 1)
        t = doc.add_table(rows=0, cols=2); t.style = "Light List Accent 1"
        for h in data["history"]:
            r = t.add_row().cells
            r[0].text, r[1].text = str(h.get("year", "")), str(h.get("event", ""))
            r[0].width, r[1].width = Cm(3), Cm(14)

    def plain(text):
        return str(text or "").replace("**", "")

    # ---- 무엇이 다른가
    cmp_ = data["comparison"]
    if cmp_:
        doc.add_page_break()
        heading(cmp_.get("title") or f"{c['name']}, 무엇이 다른가?", 1)
        if cmp_.get("quote"):
            p = doc.add_paragraph(); r = p.add_run(f"“{plain(cmp_['quote'])}”"); r.bold = True; r.font.size = Pt(14); r.font.color.rgb = brand
        if cmp_.get("before") or cmp_.get("after"):
            t = doc.add_table(rows=1, cols=2); t.style = "Light Grid Accent 1"
            t.rows[0].cells[0].text = cmp_.get("before_title") or "기존 교육의 한계"
            t.rows[0].cells[1].text = cmp_.get("after_title") or c["name"]
            n = max(len(cmp_.get("before") or []), len(cmp_.get("after") or []))
            for i in range(n):
                cells = t.add_row().cells
                b = (cmp_.get("before") or [])[i:i+1]; a = (cmp_.get("after") or [])[i:i+1]
                cells[0].text = f"✕ {b[0]}" if b else ""
                cells[1].text = f"✔ {a[0]}" if a else ""
        if cmp_.get("goal"):
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run("교육의 목표  ").font.color.rgb = RGBColor(0x88, 0x88, 0x88)
            r = p.add_run(cmp_["goal"]); r.bold = True; r.font.size = Pt(13)

    # ---- 교육 3축
    if data["program_axes"]:
        heading("교육 프로그램", 1)
        for a in data["program_axes"]:
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(f"{a.get('title', '')}  ").bold = True
            p.add_run(plain(a.get("text")))
        if data["program_axes_note"]:
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run(f"“{data['program_axes_note']}”").bold = True

    # ---- 프로그램
    if data["programs"]:
        doc.add_page_break()
        ph = data["programs_heading"]
        heading(ph.get("title") or "교육 프로그램", 1)
        if ph.get("subtitle"):
            doc.add_paragraph(ph["subtitle"])
        for p_ in data["programs"]:
            heading(p_.get("title", ""), 2)
            kv_table([("활용 대상", p_.get("target")), ("시간", p_.get("duration"))])
            if p_.get("summary"):
                doc.add_paragraph(plain(p_["summary"]))
            bullets(p_.get("features"))
            for i, sec in enumerate(p_.get("sections") or [], 1):
                heading(f"{i}. {sec.get('title', '')}", 3)
                if sec.get("subtitle"):
                    doc.add_paragraph(sec["subtitle"])
                bullets(sec.get("items"))
            for photo in p_.get("photo_paths") or []:
                if (base / photo).is_file():
                    doc.add_picture(str(base / photo), width=Cm(7))

    # ---- 강점
    if data["strengths"]:
        heading("차별화 포인트", 1)
        for i, s in enumerate(data["strengths"], 1):
            heading(f"{i}. {s.get('title', '')}", 3)
            if s.get("description"):
                doc.add_paragraph(s["description"])

    # ---- 강사
    if data["instructors"]:
        doc.add_page_break()
        heading("강사 소개", 1)
        for ins in data["instructors"]:
            t = doc.add_table(rows=1, cols=2)
            left, right = t.rows[0].cells
            left.width, right.width = Cm(4.5), Cm(12.5)
            photo = ins.get("photo_path")
            if photo and (base / photo).is_file():
                left.paragraphs[0].add_run().add_picture(str(base / photo), width=Cm(4))
            else:
                left.text = ""
            rp = right.paragraphs[0]
            r = rp.add_run(ins.get("name", "")); r.bold = True; r.font.size = Pt(15); r.font.color.rgb = brand
            if ins.get("title"):
                rp.add_run(f"   {ins['title']}").font.size = Pt(10)
            if ins.get("specialties"):
                right.add_paragraph("전문 분야: " + " · ".join(ins["specialties"]))
            if ins.get("quote"):
                q = right.add_paragraph(); q.add_run(f"“{ins['quote']}”").italic = True
            for para in ins.get("bio", []):
                right.add_paragraph(para)
            for label, key in (("주요 경력", "career"), ("학력", "education"), ("자격", "certifications")):
                if ins.get(key):
                    right.add_paragraph().add_run(label).bold = True
                    for x in ins[key]:
                        right.add_paragraph(str(x), style="List Bullet")
            doc.add_paragraph()

    # ---- 고객사 / 후기 / 마무리
    if data["clients"] or data["testimonials"] or data["closing"]:
        doc.add_page_break()
    if data["clients"]:
        heading("주요 고객사", 1)
        doc.add_paragraph("  ·  ".join(str(x) for x in data["clients"]))
    if data["testimonials"]:
        heading("고객 후기", 1)
        for t_ in data["testimonials"]:
            p = doc.add_paragraph(); p.add_run(f"“{t_.get('quote', '')}”").italic = True
            if t_.get("author"):
                doc.add_paragraph(f"— {t_['author']}").paragraph_format.left_indent = Cm(1)
    if data["closing"]:
        doc.add_paragraph()
        if data["closing"].get("headline"):
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(data["closing"]["headline"]); r.bold = True; r.font.size = Pt(16); r.font.color.rgb = brand
        if data["closing"].get("message"):
            p = doc.add_paragraph(data["closing"]["message"]); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading("문의", 2)
    kv_table([("전화", c.get("phone")), ("이메일", c.get("email")), ("웹사이트", c.get("website")),
              ("주소", c.get("address"))])

    doc.save(str(out))
    return out


# --------------------------------------------------------------------------- #
def main() -> int:
    ap = argparse.ArgumentParser(description="회사 소개서 자동 생성기")
    ap.add_argument("-d", "--data", type=Path, default=DEFAULT_DATA, help="입력 YAML (기본: data/company.yaml)")
    ap.add_argument("-o", "--out", type=Path, default=DEFAULT_OUT, help="출력 폴더 (기본: output/)")
    ap.add_argument("-n", "--name", default="company-introduction", help="출력 파일 이름 (확장자 제외)")
    ap.add_argument("--pdf", action="store_true", help="PDF 도 생성 (Chrome/Chromium 필요)")
    ap.add_argument("--docx", action="store_true", help="DOCX 도 생성 (python-docx 필요)")
    ap.add_argument("--all", action="store_true", help="HTML + PDF + DOCX 모두 생성")
    args = ap.parse_args()

    if args.all:
        args.pdf = args.docx = True

    base = args.data.resolve().parent.parent if args.data.resolve().parent.name == "data" else ROOT
    data = normalize(load_data(args.data), base)
    args.out.mkdir(parents=True, exist_ok=True)

    made = []
    html_path = build_html(data, args.out / f"{args.name}.html")
    made.append(html_path)
    if args.pdf:
        p = build_pdf(html_path, args.out / f"{args.name}.pdf")
        if p:
            made.append(p)
    if args.docx:
        p = build_docx(data, args.out / f"{args.name}.docx", base)
        if p:
            made.append(p)

    print("생성 완료:")
    for m in made:
        print(f"  - {m.relative_to(ROOT) if m.is_relative_to(ROOT) else m}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
